"""When Docling is unavailable, DeepDoc must still parse a real DOCX locally.

Remote CI has Docling cached, so ``check_installation()`` is True and the
python-docx fallback (and the old-format table translator it feeds) never
runs. These tests pin that contract at ``DeepDocParser.parse`` without
depending on ambient Docling state: the Docling parser is a runtime seam
that reports unavailable and trips if ``parse_docx`` is still called.
"""

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from xagent.providers.pdf_parser import deepdoc as deepdoc_module
from xagent.providers.pdf_parser.deepdoc import DeepDocParser

REMOTE_URL_ENV = "XAGENT_DEEPDOC_XINFERENCE_URL"


@pytest.fixture(autouse=True)
def remote_env_unset(monkeypatch: pytest.MonkeyPatch) -> None:
    """Start every test from "remote not configured"."""
    monkeypatch.delenv(REMOTE_URL_ENV, raising=False)


def _write_enrollment_docx(path: Path) -> None:
    """Write a DOCX whose fallback output is known independently of DeepDoc.

    The empty paragraph, whitespace-only cell, and blank table row exist so
    the consumer result can prove those inputs are dropped rather than
    emitted as empty segments or `` | `` cells.
    """
    document = Document()
    document.add_heading("Enrollment Review", level=1)
    document.add_paragraph("")
    document.add_paragraph("Direct entry track.")

    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Track"
    table.cell(0, 1).text = "Candidate"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "Direct Entry"
    table.cell(1, 1).text = "A-0001"
    table.cell(1, 2).text = "  "

    blank_row = table.add_row()
    for cell in blank_row.cells:
        cell.text = "   "

    document.save(path)


class UnavailableDoclingParser(deepdoc_module.DeepDocDoclingParser):  # type: ignore[misc]
    """Docling parser stand-in that reports the runtime is missing."""

    def __init__(self) -> None:
        pass

    def check_installation(self) -> bool:
        return False

    def parse_docx(self, file_path: Any, **kwargs: Any) -> Any:
        raise AssertionError(
            "parse_docx must not run when Docling reports it is unavailable"
        )


def _arm_unavailable_docling(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        DeepDocParser,
        "_get_parser_for_ext",
        lambda self, ext: UnavailableDoclingParser(),
    )


@pytest.mark.asyncio
async def test_docx_parse_falls_back_when_docling_is_unavailable(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A real DOCX still yields heading, body, and table rows without Docling."""
    docx_file = tmp_path / "enrollment.docx"
    _write_enrollment_docx(docx_file)
    _arm_unavailable_docling(monkeypatch)

    result = await DeepDocParser().parse(str(docx_file), doc_id="enrollment-doc")

    assert [
        (segment.text, segment.metadata.get("style"))
        for segment in result.text_segments
    ] == [
        ("Enrollment Review", "Heading 1"),
        ("Direct entry track.", "Normal"),
    ]
    assert all(
        segment.metadata["deepdoc_backend"] == "local"
        for segment in result.text_segments
    )
    assert len(result.tables) == 1
    table = result.tables[0]
    assert table.html == "Track | Candidate | Status\nDirect Entry | A-0001"
    assert table.metadata["type"] == "table"
    assert table.metadata["parser"] == "deepdoc"
    assert table.metadata["doc_id"] == "enrollment-doc"
    assert table.metadata["deepdoc_backend"] == "local"
    assert not result.figures
